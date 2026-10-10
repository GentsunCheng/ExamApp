import os
import json
import base64
import re
from abc import ABC, abstractmethod

import openai
import anthropic

from database import get_setting, set_setting
from crypto_util import encrypt_text, decrypt_text
from language import tr

# ---------- 常量 ----------
PROVIDER_OPENAI = 'openai'
PROVIDER_ANTHROPIC = 'anthropic'

DEFAULT_BASE_URLS = {
    PROVIDER_OPENAI: 'https://api.openai.com/v1',
    PROVIDER_ANTHROPIC: 'https://api.anthropic.com',
}

# 支持文件(文档/图片)输入的模型白名单规则
FILE_CAPABLE_PATTERNS = {
    PROVIDER_OPENAI: [
        r'^gpt-4o', r'^gpt-4\.1', r'^gpt-5', r'^o1', r'^o3', r'^o4',
    ],
    PROVIDER_ANTHROPIC: [
        r'^claude-3', r'^claude-4',
    ],
}

MAX_FILES = 5
MAX_PDF_NATIVE_SIZE = 20 * 1024 * 1024  # 20MB 以内原生上传，超出降级为文本提取
MAX_EXTRACT_CHARS = 50000

VALID_TYPES = ('single', 'multiple', 'truefalse', 'fill', 'essay')

SYSTEM_PROMPT = """# 角色
你是一名严谨的考试命题专家，擅长根据给定的学习资料命制高质量试题。

# 任务
根据用户提供的资料内容，严格按照要求的题型与数量命制试题。

# 命题要求
1. 所有题目必须完全基于提供的资料内容命制，不得编造资料中不存在的知识点。
2. 题干表述清晰、无歧义；出题语言与资料语言保持一致。
3. 各题型规则：
   - 单选题：4 个选项，仅 1 个正确答案，干扰项需有迷惑性但明确错误。
   - 多选题：4~6 个选项，正确答案不少于 2 个。
   - 判断题：无选项，陈述必须能明确判断对错，避免模棱两可。
   - 填空题：题干中用"____"标示空缺处；答案需列出所有可接受的写法（含同义表述、缩写与全称）。
   - 简答题：提供完整参考答案与评分要点。
4. 题目之间不得重复或高度相似，知识点应尽量覆盖资料的不同章节/要点。
5. 难度适中，考察对资料内容的理解与应用，而非死记硬背冷门细节。

# 输出格式（必须严格遵守）
只输出一个 JSON 对象，禁止输出任何解释、注释或 markdown 代码块标记：
{
  "questions": [
    {
      "type": "single | multiple | truefalse | fill | essay",
      "text": "题干内容",
      "options": [{"key": "A", "text": "选项内容"}],
      "correct": <见各题型规则>,
      "score": <用户指定的每题分值>
    }
  ]
}

各题型字段规则：
- single：options 为 A~D 共 4 项；correct 为单元素字母数组，如 ["A"]
- multiple：options 为 4~6 项；correct 为字母数组且长度 ≥ 2，如 ["A","C"]
- truefalse：options 为 []；correct 为布尔数组，如 [true] 或 [false]
- fill：options 为 []；correct 为字符串数组，如 ["HTTP", "超文本传输协议"]
- essay：options 为 []；correct 为单元素字符串数组，内容为参考答案与评分要点
"""

RETRY_PROMPT_TEMPLATE = "你上次的输出不符合 JSON 格式要求：{error}。请重新输出，只输出符合规范的 JSON 对象。"


# ---------- 配置存取 ----------
def load_ai_config():
    def dec(v):
        if not v:
            return ''
        try:
            return decrypt_text(v) or ''
        except Exception:
            return ''
    return {
        'provider': get_setting('ai_provider') or PROVIDER_OPENAI,
        'base_url_openai': get_setting('ai_base_url_openai') or DEFAULT_BASE_URLS[PROVIDER_OPENAI],
        'base_url_anthropic': get_setting('ai_base_url_anthropic') or DEFAULT_BASE_URLS[PROVIDER_ANTHROPIC],
        'api_key_openai': dec(get_setting('ai_api_key_openai')),
        'api_key_anthropic': dec(get_setting('ai_api_key_anthropic')),
        'model': get_setting('ai_model') or '',
    }


def save_ai_config(cfg):
    set_setting('ai_provider', cfg.get('provider') or PROVIDER_OPENAI)
    set_setting('ai_base_url_openai', cfg.get('base_url_openai') or DEFAULT_BASE_URLS[PROVIDER_OPENAI])
    set_setting('ai_base_url_anthropic', cfg.get('base_url_anthropic') or DEFAULT_BASE_URLS[PROVIDER_ANTHROPIC])
    if cfg.get('api_key_openai'):
        set_setting('ai_api_key_openai', encrypt_text(cfg['api_key_openai']))
    if cfg.get('api_key_anthropic'):
        set_setting('ai_api_key_anthropic', encrypt_text(cfg['api_key_anthropic']))
    set_setting('ai_model', cfg.get('model') or '')


# ---------- 文件处理 ----------
def _extract_docx_text(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts)


def _extract_pptx_text(path):
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = ''.join(run.text for run in para.runs).strip()
                    if t:
                        slide_texts.append(t)
        if slide_texts:
            parts.append(f'[第{i}页]\n' + '\n'.join(slide_texts))
    return '\n'.join(parts)


def _extract_pdf_text(path):
    import fitz
    doc = fitz.open(path)
    parts = []
    for page in doc:
        t = page.get_text()
        if t.strip():
            parts.append(t)
    doc.close()
    return '\n'.join(parts)


def prepare_file(path):
    """返回文件描述结构:
    {'kind': 'native_pdf'|'native_image'|'text', 'path', 'name', 'mime', 'data'(base64) 或 'text'}
    """
    ext = os.path.splitext(path)[1].lower()
    name = os.path.basename(path)
    size = os.path.getsize(path)
    if ext == '.pdf':
        if size <= MAX_PDF_NATIVE_SIZE:
            with open(path, 'rb') as f:
                data = base64.b64encode(f.read()).decode('ascii')
            return {'kind': 'native_pdf', 'path': path, 'name': name,
                    'mime': 'application/pdf', 'data': data}
        text = _extract_pdf_text(path)
        return {'kind': 'text', 'path': path, 'name': name, 'text': text[:MAX_EXTRACT_CHARS]}
    if ext in ('.png', '.jpg', '.jpeg', '.webp', '.gif'):
        mime = {'.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg',
                '.webp': 'image/webp', '.gif': 'image/gif'}[ext]
        with open(path, 'rb') as f:
            data = base64.b64encode(f.read()).decode('ascii')
        return {'kind': 'native_image', 'path': path, 'name': name, 'mime': mime, 'data': data}
    if ext == '.docx':
        text = _extract_docx_text(path)
    elif ext == '.pptx':
        text = _extract_pptx_text(path)
    elif ext in ('.txt', '.md', '.csv', '.json', '.yaml', '.yml'):
        with open(path, 'r', encoding='utf-8', errors='ignore') as f:
            text = f.read()
    else:
        raise ValueError(tr('ai.error.unsupported_file', ext=ext))
    return {'kind': 'text', 'path': path, 'name': name, 'text': text[:MAX_EXTRACT_CHARS]}


# ---------- Provider 抽象 ----------
class AIProvider(ABC):
    def __init__(self, base_url, api_key):
        self.base_url = base_url
        self.api_key = api_key

    @abstractmethod
    def list_models(self):
        """返回模型 id 列表"""

    @abstractmethod
    def _build_messages(self, system_prompt, user_prompt, files):
        """构造各自 SDK 的消息结构，返回可直接传给 SDK 的参数"""

    @abstractmethod
    def _call(self, model, built):
        """执行调用，返回文本内容"""

    def generate(self, model, system_prompt, user_prompt, files):
        built = self._build_messages(system_prompt, user_prompt, files)
        return self._call(model, built)

    def filter_file_capable(self, model_ids):
        patterns = FILE_CAPABLE_PATTERNS.get(self.provider_name(), [])
        out = []
        for mid in model_ids:
            if any(re.search(p, mid) for p in patterns):
                out.append(mid)
        return out

    @classmethod
    def provider_name(cls):
        raise NotImplementedError


class OpenAIProvider(AIProvider):
    @classmethod
    def provider_name(cls):
        return PROVIDER_OPENAI

    def _client(self):
        return openai.OpenAI(base_url=self.base_url, api_key=self.api_key, timeout=300)

    def list_models(self):
        client = self._client()
        resp = client.models.list()
        ids = []
        for m in resp.data:
            mid = getattr(m, 'id', None)
            if mid:
                ids.append(mid)
        return sorted(ids)

    def _build_messages(self, system_prompt, user_prompt, files):
        content = []
        for f in files:
            if f['kind'] in ('native_pdf', 'native_image'):
                url = f"data:{f['mime']};base64,{f['data']}"
                content.append({'type': 'image_url', 'image_url': {'url': url}})
        content.append({'type': 'text', 'text': user_prompt})
        return {
            'messages': [
                {'role': 'system', 'content': system_prompt},
                {'role': 'user', 'content': content},
            ]
        }

    def _call(self, model, built):
        client = self._client()
        resp = client.chat.completions.create(
            model=model,
            messages=built['messages'],
            temperature=0.7,
        )
        return resp.choices[0].message.content or ''


class AnthropicProvider(AIProvider):
    @classmethod
    def provider_name(cls):
        return PROVIDER_ANTHROPIC

    def _client(self):
        return anthropic.Anthropic(base_url=self.base_url, api_key=self.api_key, timeout=300)

    def list_models(self):
        client = self._client()
        resp = client.models.list()
        ids = []
        for m in resp.data:
            mid = getattr(m, 'id', None)
            if mid:
                ids.append(mid)
        return sorted(ids)

    def _build_messages(self, system_prompt, user_prompt, files):
        content = []
        for f in files:
            if f['kind'] == 'native_pdf':
                content.append({
                    'type': 'document',
                    'source': {
                        'type': 'base64',
                        'media_type': 'application/pdf',
                        'data': f['data'],
                    },
                })
            elif f['kind'] == 'native_image':
                content.append({
                    'type': 'image',
                    'source': {
                        'type': 'base64',
                        'media_type': f['mime'],
                        'data': f['data'],
                    },
                })
        content.append({'type': 'text', 'text': user_prompt})
        return {
            'system': system_prompt,
            'messages': [{'role': 'user', 'content': content}],
        }

    def _call(self, model, built):
        client = self._client()
        resp = client.messages.create(
            model=model,
            max_tokens=8192,
            system=built['system'],
            messages=built['messages'],
            temperature=0.7,
        )
        parts = []
        for block in resp.content:
            if getattr(block, 'type', None) == 'text':
                parts.append(block.text)
        return '\n'.join(parts)


def create_provider(provider_name, base_url, api_key):
    if provider_name == PROVIDER_ANTHROPIC:
        return AnthropicProvider(base_url, api_key)
    return OpenAIProvider(base_url, api_key)


# ---------- 提示词构建 ----------
def build_user_prompt(counts, score, materials_text):
    """counts: {'single': n, 'multiple': n, ...}，仅包含数量>0 的题型"""
    type_names = {
        'single': '单选题', 'multiple': '多选题', 'truefalse': '判断题',
        'fill': '填空题', 'essay': '简答题',
    }
    req_parts = [f"{type_names[t]} {n} 道" for t, n in counts.items() if n > 0]
    req = '、'.join(req_parts) + f"，每题 {score:g} 分"
    prompt = f"# 数量要求\n{req}\n\n# 资料内容\n{materials_text}"
    return prompt


def build_materials_text(files):
    """将 text 类型文件拼为资料文本；原生上传的文件仅列出文件名提示"""
    parts = []
    native_names = [f['name'] for f in files if f['kind'] != 'text']
    if native_names:
        parts.append('以下文件已作为附件直接提供，请阅读附件内容：' + '、'.join(native_names))
    for f in files:
        if f['kind'] == 'text':
            parts.append(f"## 资料文件：{f['name']}\n{f['text']}")
    return '\n\n'.join(parts)


# ---------- 结果解析与校验 ----------
def _strip_code_fence(raw):
    s = raw.strip()
    if s.startswith('```'):
        s = re.sub(r'^```[a-zA-Z]*\s*', '', s)
        s = re.sub(r'\s*```$', '', s)
    return s.strip()


def validate_question(q, expected_score):
    """校验单题结构，返回 (ok, error_msg)"""
    if not isinstance(q, dict):
        return False, '题目不是对象'
    qtype = q.get('type')
    if qtype not in VALID_TYPES:
        return False, f'非法题型: {qtype}'
    text = q.get('text')
    if not isinstance(text, str) or not text.strip():
        return False, '题干为空'
    options = q.get('options')
    correct = q.get('correct')
    if not isinstance(options, list):
        return False, 'options 不是数组'
    if not isinstance(correct, list) or len(correct) == 0:
        return False, 'correct 为空或不是数组'

    if qtype in ('single', 'multiple'):
        keys = []
        for opt in options:
            if not isinstance(opt, dict) or 'key' not in opt or 'text' not in opt:
                return False, '选项结构错误，需含 key 与 text'
            keys.append(str(opt['key']).upper())
        if qtype == 'single' and len(options) != 4:
            return False, f'单选题需 4 个选项，实际 {len(options)}'
        if qtype == 'multiple' and not (4 <= len(options) <= 6):
            return False, f'多选题需 4~6 个选项，实际 {len(options)}'
        letters = [str(c).upper() for c in correct]
        if qtype == 'single' and len(letters) != 1:
            return False, '单选题正确答案必须为 1 个'
        if qtype == 'multiple' and len(letters) < 2:
            return False, '多选题正确答案不少于 2 个'
        for lt in letters:
            if lt not in keys:
                return False, f'正确答案 {lt} 不在选项内'
        q['correct'] = letters
        q['options'] = [{'key': str(o['key']).upper(), 'text': str(o['text'])} for o in options]
    elif qtype == 'truefalse':
        if options:
            return False, '判断题不应有选项'
        v = correct[0]
        if isinstance(v, str):
            lv = v.strip().lower()
            if lv in ('true', '对', '正确', 'yes'):
                v = True
            elif lv in ('false', '错', '错误', 'no'):
                v = False
            else:
                return False, '判断题答案必须为 true/false'
        if not isinstance(v, bool):
            return False, '判断题答案必须为布尔值'
        q['correct'] = [v]
        q['options'] = []
    elif qtype == 'fill':
        if options:
            return False, '填空题不应有选项'
        answers = [str(a).strip() for a in correct if str(a).strip()]
        if not answers:
            return False, '填空题答案为空'
        q['correct'] = answers
        q['options'] = []
    elif qtype == 'essay':
        if options:
            return False, '简答题不应有选项'
        ref = str(correct[0]).strip()
        if not ref:
            return False, '简答题参考答案为空'
        q['correct'] = [ref]
        q['options'] = []

    sc = q.get('score', expected_score)
    try:
        q['score'] = float(sc)
    except Exception:
        q['score'] = float(expected_score)
    return True, ''


def parse_and_validate(raw_text, expected_score):
    """解析模型输出，返回 (questions, error)。questions 为校验通过的题目列表"""
    s = _strip_code_fence(raw_text)
    try:
        obj = json.loads(s)
    except Exception as e:
        return None, f'JSON 解析失败: {e}'
    qs = obj.get('questions') if isinstance(obj, dict) else None
    if not isinstance(qs, list) or not qs:
        return None, '输出缺少 questions 数组'
    valid = []
    errors = []
    for i, q in enumerate(qs, 1):
        ok, err = validate_question(q, expected_score)
        if ok:
            valid.append(q)
        else:
            errors.append(f'第{i}题: {err}')
    if not valid:
        return None, '所有题目均校验失败：' + '；'.join(errors[:5])
    return valid, ('；'.join(errors[:5]) if errors else None)


def generate_questions(provider, model, counts, score, files, on_retry=None):
    """主入口：调用模型生成题目，解析失败自动重试一次。
    返回 (questions, error)"""
    materials = build_materials_text(files)
    user_prompt = build_user_prompt(counts, score, materials)
    try:
        raw = provider.generate(model, SYSTEM_PROMPT, user_prompt, files)
    except Exception as e:
        return None, f'模型调用失败: {e}'
    questions, err = parse_and_validate(raw, score)
    if questions is not None and err is None:
        return questions, None
    # 重试一次
    if on_retry:
        on_retry(err or '格式错误')
    retry_prompt = user_prompt + '\n\n' + RETRY_PROMPT_TEMPLATE.format(error=err or '格式错误')
    try:
        raw2 = provider.generate(model, SYSTEM_PROMPT, retry_prompt, files)
    except Exception as e:
        return None, f'重试调用失败: {e}'
    questions2, err2 = parse_and_validate(raw2, score)
    if questions2:
        return questions2, err2
    return None, err2 or err or '生成失败'
